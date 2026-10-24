"""Generate a local Excel snapshot in memory; no file or network upload."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from billing.person import personalized_bill

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CURRENCY = '"$"#,##0.00;[Red]("$"#,##0.00)'


def _append(sheet, values):
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in values])


def workbook_bytes(person=None, state=None, kb=None):
    """Export saved local details and an optional completed review snapshot.

    Free text is explicitly stored as text: never formulas or hyperlinks.
    No raw document, conversation, account credential or remote call is used.
    """
    person = person or {}
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    _append(summary, ["ClearClaim medical bill review"])
    _append(summary, ["Educational snapshot • all reference prices and policies are fictional"])
    _append(summary, ["Export date", date.today()])
    _append(summary, ["Personal details", "Saved locally for this session"])
    _append(summary, ["Name", person.get("name", "")])
    _append(summary, ["Age", person.get("age")])
    _append(summary, ["Sex", person.get("sex", "Not provided")])
    _append(summary, ["Bill reference", person.get("reference", "")])
    _append(summary, ["Age and sex are not used by the audit and are not sent to any model."])
    _append(summary, [])
    if state and "audit" in state and "bill" in state:
        bill = personalized_bill(state["bill"], person)
        audit = state["audit"]
        for row in [
            ["Confirmed bill", "Value"], ["Patient on exported bill", bill["patient"]],
            ["Provider", bill["provider"]], ["Account", bill["account"]],
            ["Total billed (USD)", Decimal(audit["total_billed"])],
            ["Illustrative potential reduction (USD)", Decimal(audit["potential_savings"])],
            ["Illustrative remaining amount (USD)", Decimal(audit["illustrative_remaining"])],
            ["Rule findings", len(audit["findings"])],
            ["Processing", "Fixed synthetic demo" if state.get("cloud_eligible") else "Local only"],
            ["These are snapshot totals, not guarantees or editable reimbursement calculations."],
            ["Overlapping amounts are counted once per line; do not sum standalone finding amounts."],
        ]:
            _append(summary, row)
        for row_number in (15, 16, 17):
            summary.cell(row_number, 2).number_format = CURRENCY
        charges = book.create_sheet("Charges")
        _append(charges, ["Confirmed charges • local export"])
        _append(charges, ["Source: confirmed review table; billed amounts are line totals in USD"])
        _append(charges, ["Line ID", "Procedure code", "Description", "Service date", "Quantity",
                        "Billed total (USD)", "Modifiers", "Disputed amount (USD)"])
        for line in bill["lines"]:
            day = date.fromisoformat(line["service_date"]) if line.get("service_date") else None
            _append(charges, [line["line_id"], line["code"], line["description"], day, line["quantity"],
                            Decimal(str(line["amount"])), line.get("modifiers", ""),
                            Decimal(audit["disputed_by_line"].get(line["line_id"], "0"))])
        for row in charges.iter_rows(min_row=4):
            row[3].number_format = "yyyy-mm-dd"
            row[5].number_format = CURRENCY
            row[7].number_format = CURRENCY
        findings = book.create_sheet("Findings")
        _append(findings, ["Possible billing issues • not proof of fraud"])
        _append(findings, ["Standalone candidate amounts can overlap; use the Summary reduction total"])
        _append(findings, ["Finding", "Category", "Title", "Line IDs", "Severity", "Confidence",
                         "Explanation", "Standalone amount (USD)", "Reference IDs"])
        for item in audit["findings"]:
            _append(findings, [item["id"], item["category"], item["title"], ", ".join(item["line_ids"]),
                             item["severity"], item["confidence"], item["detail"],
                             Decimal(item["candidate_amount"]), ", ".join(item["citations"])])
            findings.cell(findings.max_row, 8).number_format = CURRENCY
        if not audit["findings"]:
            _append(findings, ["No findings matched the limited fictional rules. This does not certify the bill as correct."])
        if kb is not None:
            refs = book.create_sheet("References")
            _append(refs, ["Fictional reference evidence"])
            _append(refs, ["Source: bundled synthetic policies/pricing; not Medicare or insurer rates"])
            _append(refs, ["Reference ID", "Title", "Evidence text"])
            ids = {sid for item in audit["findings"] for sid in item["citations"]} | {"POL-SAVINGS", "POL-DISPUTE"}
            for sid in sorted(ids):
                document = kb.get(sid)
                if document:
                    _append(refs, [sid, document.metadata["title"], document.page_content])
    else:
        _append(summary, ["No completed bill review is included. Analyze a bill to export charges and findings."])

    if state and state.get("review") is not None:
        import json
        review_sheet = book.create_sheet("Review decisions")
        _append(review_sheet, ["User review snapshot"])
        _append(review_sheet, ["Original audit findings remain unchanged"])
        _append(review_sheet, ["Section", "Data"])
        for key, value in state["review"].items():
            entries = value if isinstance(value, list) else [value]
            for entry in entries:
                text = json.dumps(entry, ensure_ascii=False)
                for offset in range(0, len(text), 30000):
                    _append(review_sheet, [key, text[offset:offset + 30000]])
        if state.get("references"):
            _append(review_sheet, ["Reference versions", json.dumps(state["references"])])
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A4" if sheet.title != "Summary" else "B5"
        sheet.row_dimensions[1].height = 30
        sheet.row_dimensions[2].height = 32
        for row in sheet:
            for cell in row:
                # Explicit inline strings prevent Excel formula / external-link injection.
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="183D35")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == 1 or (cell.row == 3 and sheet.title != "Summary"):
                    cell.fill = PatternFill("solid", fgColor="144F43")
                    cell.font = Font(name="Calibri", size=14 if cell.row == 1 else 11,
                                     bold=True, color="FFFFFF")
                elif cell.row > 3 and cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F6F3")
                if sheet.title == "Summary" and cell.row in (4, 11):
                    cell.fill = PatternFill("solid", fgColor="E3F0E8")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="144F43")
        if sheet.title != "Summary":
            sheet.auto_filter.ref = f"A3:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        widths = {"Summary": [52, 65], "Charges": [15, 18, 42, 17, 12, 23, 18, 25],
                  "Findings": [14, 22, 40, 23, 15, 42, 90, 26, 38], "References": [23, 42, 110]}
        for i, width in enumerate(widths.get(sheet.title, [25] * sheet.max_column), 1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        # Fit wrapped explanations without shrinking their font.
        for row in sheet.iter_rows(min_row=3):
            lines = max((len(str(cell.value or "")) // max(10, int(sheet.column_dimensions[cell.column_letter].width) - 3) + 1
                         for cell in row), default=1)
            sheet.row_dimensions[row[0].row].height = min(200, max(24, lines * 16))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
